#!/usr/bin/env python3
"""Estimate man-days (MD) from a requirement list and write an Excel summary.

Usage: estimate.py REQS.(xlsx|csv|md) [--out Estimate.xlsx] [--qa 25] [--pm 15] [--contingency 15]
       [--rate 0]   (rate = THB per MD; 0 = no cost columns)
Input columns (case-insensitive): REQ ID, Requirement, Fit, Complexity [, Effort (MD)]
Fit: Standard | Config | Studio | Custom | Report | Integration | Migration | Out of scope
Complexity: S | M | L | XL.  An Effort (MD) value overrides the table.
Baseline table = references/estimation.md; adjust to your team's history.
"""
import argparse
import csv
import os
import re
import sys

BASE = {  # fit -> (S, M, L, XL) man-days
    "standard": (0.25, 0.5, 1, 2),
    "config": (0.5, 1, 2, 4),
    "studio": (0.5, 1, 2, 4),
    "custom": (1, 3, 6, 12),
    "report": (1, 2, 4, 6),
    "integration": (3, 6, 12, 20),
    "migration": (1, 2, 4, 8),
    "out of scope": (0, 0, 0, 0),
}
CX = {"s": 0, "m": 1, "l": 2, "xl": 3}
BUILD = {"config", "studio", "custom", "report", "integration", "migration"}


def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    elif ext == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = [[c.strip() for c in r] for r in csv.reader(f)]
    else:  # markdown: first table containing a 'fit' header
        rows, cur = [], []
        for line in open(path, encoding="utf-8"):
            if line.strip().startswith("|"):
                if re.fullmatch(r"\|?(\s*:?-{2,}:?\s*\|)+\s*:?-*:?\s*\|?", line.strip()):
                    continue
                cur.append([c.strip() for c in line.strip().strip("|").split("|")])
            elif cur:
                if any(h.lower() == "fit" for h in cur[0]):
                    rows = cur
                    break
                cur = []
        if not rows and cur and any(h.lower() == "fit" for h in cur[0]):
            rows = cur
    rows = [r for r in rows if any(r)]
    if not rows:
        sys.exit("no rows found")
    hdr = [h.lower() for h in rows[0]]
    return hdr, rows[1:]


def col(hdr, *names):
    for n in names:
        if n in hdr:
            return hdr.index(n)
    return None


def main():
    p = argparse.ArgumentParser()
    p.add_argument("reqs")
    p.add_argument("--out", default=None)
    p.add_argument("--qa", type=float, default=25)
    p.add_argument("--pm", type=float, default=15)
    p.add_argument("--contingency", type=float, default=15)
    p.add_argument("--rate", type=float, default=0)
    a = p.parse_args()
    hdr, rows = read_rows(a.reqs)
    i_id, i_req = col(hdr, "req id", "id", "req"), col(hdr, "requirement", "description")
    i_fit, i_cx = col(hdr, "fit", "fit/gap", "classification"), col(hdr, "complexity", "size")
    i_eff = col(hdr, "effort (md)", "effort", "md")
    if i_fit is None:
        sys.exit("missing column: Fit")
    lines, problems = [], []
    for n, r in enumerate(rows, 2):
        g = lambda i: r[i] if i is not None and i < len(r) else ""  # noqa: E731
        fit, cx = g(i_fit).lower(), (g(i_cx) or "M").lower()
        override = g(i_eff).replace(",", "")
        if fit not in BASE:
            problems.append(f"row {n}: unknown Fit '{g(i_fit)}'")
            continue
        if cx not in CX:
            problems.append(f"row {n}: unknown Complexity '{g(i_cx)}' (use S/M/L/XL)")
            continue
        try:
            md = float(override) if override else BASE[fit][CX[cx]]
        except ValueError:
            problems.append(f"row {n}: Effort not a number '{override}'")
            continue
        lines.append([g(i_id), g(i_req), fit.title(), cx.upper(), md, "override" if override else "table"])

    build = sum(l[4] for l in lines if l[2].lower() in BUILD)
    base = sum(l[4] for l in lines)
    qa = build * a.qa / 100
    sub = base + qa
    pm = sub * a.pm / 100
    cont = (sub + pm) * a.contingency / 100
    total = sub + pm + cont
    summary = [["Requirements (base)", base], [f"QA / testing ({a.qa:g}% of build)", qa],
               [f"Project management ({a.pm:g}%)", pm], [f"Contingency ({a.contingency:g}%)", cont],
               ["TOTAL (MD)", total]]

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Lines"
    head = ["REQ ID", "Requirement", "Fit", "Complexity", "MD", "Source"] + (["Cost (THB)"] if a.rate else [])
    ws.append(head)
    for l in lines:
        ws.append(l + ([l[4] * a.rate] if a.rate else []))
    sm = wb.create_sheet("Summary")
    sm.append(["Item", "MD"] + (["Cost (THB)"] if a.rate else []))
    for s in summary:
        sm.append(s + ([s[1] * a.rate] if a.rate else []))
    sm.append([])
    sm.append(["By Fit", "MD"])
    for fit in BASE:
        v = sum(l[4] for l in lines if l[2].lower() == fit)
        if v:
            sm.append([fit.title(), v])
    for w in (ws, sm):
        for c in w[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="714B67")
        w.column_dimensions["A"].width = 34
        w.column_dimensions["B"].width = 50 if w is ws else 14
    out = a.out or os.path.splitext(a.reqs)[0] + "_estimate.xlsx"
    wb.save(out)
    for s in summary:
        print(f"{s[0]}: {s[1]:,.2f} MD" + (f" = {s[1] * a.rate:,.0f} THB" if a.rate else ""))
    for pr in problems:
        print("SKIPPED", pr)
    print(out)


if __name__ == "__main__":
    main()
